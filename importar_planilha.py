#!/usr/bin/env python3
"""Importa a planilha 'Controle Financeiro.xlsx' para o banco do sistema (dados.db).

Uso (precisa do openpyxl só para esta importação):
    pip install openpyxl
    python3 importar_planilha.py "../Controle Financeiro.xlsx"

Apaga e recria dados.db. Traz: categorias, recorrências, parcelamentos, compras
avulsas de cada mês, entradas e as marcações de "Pago" de cada aba de mês.
"""
import re
import sys
from datetime import date, datetime

import openpyxl

import app

ABAS_MES = ["01-Janeiro", "02-Fevereiro", "03-Março", "04-Abril", "05-Maio", "06-Junho", "07-Julho",
            "08-Agosto", "09-Setembro", "10-Outubro", "11-Novembro", "12-Dezembro"]
PARCELA_RE = re.compile(r"^(.*) \(parcela (\d+)/(\d+)\)$")


def texto(v):
    return str(v).strip() if v is not None else ""


def como_data(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):           # serial do Excel/Sheets
        return date.fromordinal(date(1899, 12, 30).toordinal() + int(v))
    return None


def linhas(ws, inicio, cols):
    for r in range(inicio, ws.max_row + 1):
        yield r, [ws.cell(r, c).value for c in cols]


def main(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ano = int(wb["Config"]["B2"].value)
    avisos = []

    if app.DB_PATH.exists():
        app.DB_PATH.unlink()
    with app.conectar() as con:
        con.executescript(app.SCHEMA)

        # Categorias
        cat_id = {}
        for r, (nome,) in linhas(wb["Config"], 6, [1]):
            nome = texto(nome)
            if nome and nome not in cat_id:
                cat_id[nome] = con.execute("INSERT INTO categorias (nome, ordem) VALUES (?, ?)",
                                           (nome, len(cat_id))).lastrowid

        def categoria(nome, onde):
            nome = texto(nome) or "Outros"
            if nome not in cat_id:
                avisos.append(f"{onde}: categoria {nome!r} não existia, foi criada")
                cat_id[nome] = con.execute("INSERT INTO categorias (nome, ordem) VALUES (?, ?)",
                                           (nome, len(cat_id))).lastrowid
            return cat_id[nome]

        # Recorrências (na planilha valiam o ano inteiro; status Inativo some de todos os meses)
        rec_por_linha = {}
        for r, (desc, dia, valor, cat, _freq, status) in linhas(wb["Base de Recorrências"], 4, range(2, 8)):
            desc = texto(desc)
            if not desc:
                continue
            if valor in (None, ""):
                avisos.append(f"Recorrência {desc!r} sem valor: importada com R$ 0,00")
            fim = None if texto(status) != "Inativo" else f"{ano - 1}-12"
            rec_por_linha[r] = con.execute(
                """INSERT INTO recorrencias (descricao, dia, valor_centavos, categoria_id, inicio, fim)
                   VALUES (?, ?, ?, ?, ?, ?)""",
                (desc, int(dia or 1), app.centavos(valor or 0), categoria(cat, f"Recorrência {desc!r}"),
                 f"{ano}-01", fim)).lastrowid

        # Parcelamentos
        parc_chave = {}   # (descricao, n, valor_centavos) -> id
        for r, (desc, primeira, valor, n, cat) in linhas(wb["Base de Parcelamentos"], 4, range(2, 7)):
            desc = texto(desc)
            if not desc:
                continue
            d = como_data(primeira)
            if d is None:
                avisos.append(f"Parcelamento {desc!r} sem data da 1ª parcela: NÃO importado")
                continue
            if not valor or not n:
                avisos.append(f"Parcelamento {desc!r} sem valor e/ou nº de parcelas: importado como "
                              f"1x de R$ {float(valor or 0):.2f} — confira na aba Parcelamentos")
            pid = con.execute(
                """INSERT INTO parcelamentos (descricao, primeira_parcela, valor_parcela_centavos,
                   n_parcelas, categoria_id) VALUES (?, ?, ?, ?, ?)""",
                (desc, d.isoformat(), app.centavos(valor or 0), int(n or 1),
                 categoria(cat, f"Parcelamento {desc!r}"))).lastrowid
            parc_chave.setdefault((desc, int(n or 1), app.centavos(valor or 0)), pid)

        # Abas de mês: avulsos + marcações de Pago
        for m, aba in enumerate(ABAS_MES, start=1):
            ws, mes = wb[aba], f"{ano}-{m:02d}"
            marca_rec = marca_parc = None
            for r in range(1, ws.max_row + 1):
                v = texto(ws.cell(r, 2).value)
                if v.startswith("Recorrências deste mês"):
                    marca_rec = r
                elif v.startswith("Parcelas deste mês"):
                    marca_parc = r

            for r in range(4, marca_rec - 1):
                desc, dt, valor, cat, pago = (ws.cell(r, c).value for c in range(2, 7))
                if not texto(desc):
                    continue
                d = como_data(dt)
                con.execute("""INSERT INTO lancamentos (mes, descricao, data, valor_centavos, categoria_id, pago)
                               VALUES (?, ?, ?, ?, ?, ?)""",
                            (mes, texto(desc), d.isoformat() if d else None, app.centavos(valor or 0),
                             categoria(cat, f"{aba} {desc!r}"), 1 if pago else 0))

            # Recorrências: linha i do bloco aponta para a linha 4+i da base
            for i, r in enumerate(range(marca_rec + 2, marca_parc - 1)):
                if ws.cell(r, 6).value is True and (4 + i) in rec_por_linha and texto(ws.cell(r, 2).value):
                    con.execute("INSERT OR IGNORE INTO pagamentos VALUES ('rec', ?, ?)", (rec_por_linha[4 + i], mes))

            # Parcelas: identifica pelo texto "Descrição (parcela k/n)" + valor
            for r in range(marca_parc + 2, ws.max_row + 1):
                achou = PARCELA_RE.match(texto(ws.cell(r, 2).value))
                if achou and ws.cell(r, 6).value is True:
                    chave = (achou.group(1).strip(), int(achou.group(3)), app.centavos(ws.cell(r, 4).value or 0))
                    if chave in parc_chave:
                        con.execute("INSERT OR IGNORE INTO pagamentos VALUES ('parc', ?, ?)", (parc_chave[chave], mes))

        # Entradas
        ws = wb["Entradas"]
        for r in range(4, ws.max_row + 1):
            nome = texto(ws.cell(r, 1).value)
            if not nome or nome.lower() == "total":
                continue
            fid = con.execute("INSERT INTO fontes_renda (nome, ordem) VALUES (?, ?)", (nome, r)).lastrowid
            for m in range(1, 13):
                v = ws.cell(r, m + 1).value
                if isinstance(v, (int, float)) and v:
                    con.execute("INSERT INTO entradas VALUES (?, ?, ?)", (fid, f"{ano}-{m:02d}", app.centavos(v)))

        resumo = {t: con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
                  for t in ("categorias", "recorrencias", "parcelamentos", "lancamentos", "pagamentos", "entradas")}

    print("Importação concluída:", ", ".join(f"{k}={v}" for k, v in resumo.items()))
    for a in avisos:
        print("  aviso:", a)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
